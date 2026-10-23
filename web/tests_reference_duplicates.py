import io
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Barrier
from unittest import skipUnless

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.db import connection, connections
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse
from openpyxl import load_workbook

from imports.reconciliation import build_preview, publish_preview
from imports.tests.test_excel import synthetic_workbook
from inventory.catalog import DuplicateReferenceError, matching_references, save_reference
from inventory.models import AuditEvent, ImportBatch, Location, Movement, MovementLine, Reference


class ReferenceDuplicateTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username="synthetic-match-admin", role="admin")
        cls.viewer = get_user_model().objects.create_user(username="synthetic-match-viewer", role="viewer")
        cls.location = Location.objects.create(name="Local sintético Alfa")
        cls.other_location = Location.objects.create(name="Local sintético Beta")
        cls.original = Reference.objects.create(
            code="SYN-ORIGINAL", name="  Referência   Fictícia Alfa  ",
            location=cls.location, current_stock=0,
        )

    def setUp(self):
        self.client.force_login(self.admin)
        self.create_url = reverse("web:reference-create")

    def data(self, name="referência fictícia alfa", **overrides):
        return {"name": name, "location": self.other_location.pk, "sector": "",
                "box_capacity": "6", **overrides}

    def test_service_blocks_normalized_match_across_locations(self):
        with self.assertRaises(DuplicateReferenceError) as caught:
            save_reference(actor=self.admin, data={
                "name": "\tREFERÊNCIA  FICTÍCIA\nALFA ", "location": self.other_location,
            })
        self.assertEqual([ref.pk for ref in caught.exception.matches], [self.original.pk])
        self.assertEqual(Reference.objects.count(), 1)
        self.assertFalse(AuditEvent.objects.exists())

    def test_post_shows_existing_code_location_state_and_cannot_override_block(self):
        response = self.client.post(self.create_url, self.data(
            name=" REFERÊNCIA FICTÍCIA   ALFA ", confirm_duplicate="on",
        ))
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, self.original.code, status_code=400)
        self.assertContains(response, self.location.name, status_code=400)
        self.assertContains(response, "Ativa", status_code=400)
        self.assertEqual(Reference.objects.count(), 1)
        self.assertFalse(AuditEvent.objects.exists())

    def test_inactive_match_suggests_reactivation_without_changing_history_or_code(self):
        self.original.active = False
        self.original.save(update_fields=["active"])
        movement = Movement.objects.create(kind="inventory", reason="Histórico sintético original")
        line = MovementLine.objects.create(
            movement=movement, reference=self.original, quantity=0,
            reference_name=self.original.name,
        )
        movement.status = "confirmed"
        movement.save(update_fields=["status"])
        response = self.client.post(self.create_url, self.data())
        self.assertContains(response, "Inativa", status_code=400)
        self.assertContains(response, "Rever reativação da referência original", status_code=400)
        activation_url = reverse("web:reference-toggle", args=[self.original.pk])
        self.assertContains(response, activation_url, status_code=400)
        self.assertEqual(self.client.get(activation_url).status_code, 200)
        self.original.refresh_from_db()
        self.assertFalse(self.original.active)
        self.assertEqual(self.original.code, "SYN-ORIGINAL")
        self.assertEqual(Reference.objects.count(), 1)
        self.assertFalse(AuditEvent.objects.exists())
        self.assertTrue(MovementLine.objects.filter(pk=line.pk, reference=self.original).exists())
        self.assertEqual(self.client.post(activation_url, {"active": "1"}).status_code, 302)
        self.original.refresh_from_db()
        self.assertTrue(self.original.active)
        self.assertEqual(self.original.code, "SYN-ORIGINAL")
        self.assertTrue(MovementLine.objects.filter(pk=line.pk, reference=self.original).exists())
        self.assertEqual(Reference.objects.count(), 1)

    def test_multiple_historical_matches_are_all_shown_and_left_untouched(self):
        other = Reference.objects.create(
            code="SYN-HISTORICAL", name="REFERÊNCIA FICTÍCIA ALFA",
            location=self.other_location, active=False,
        )
        before = list(Reference.objects.order_by("pk").values())
        response = self.client.post(self.create_url, self.data())
        self.assertContains(response, "Situação para revisão", status_code=400)
        for ref in (self.original, other):
            self.assertContains(response, ref.code, status_code=400)
            self.assertContains(response, ref.location.name, status_code=400)
        self.assertEqual(list(Reference.objects.order_by("pk").values()), before)
        self.assertFalse(AuditEvent.objects.exists())

    def test_get_name_check_is_read_only_and_prefills_creation_form(self):
        response = self.client.get(self.create_url, {"check_name": "referência  fictícia alfa"})
        self.assertContains(response, self.original.code)
        self.assertEqual(response.context["form"].initial["name"], "referência  fictícia alfa")
        self.assertEqual(Reference.objects.count(), 1)
        self.assertFalse(AuditEvent.objects.exists())

    def test_duplicate_is_shown_even_when_other_fields_are_missing(self):
        response = self.client.post(self.create_url, {"name": "referência fictícia alfa"})
        self.assertContains(response, self.original.code, status_code=400)
        self.assertEqual(Reference.objects.count(), 1)

    def test_distinct_designation_can_be_created_and_accents_remain_significant(self):
        self.assertEqual(matching_references("Referencia Ficticia Alfa"), [])
        response = self.client.post(self.create_url, self.data(name="Referência fictícia Beta"))
        self.assertEqual(response.status_code, 302)
        created = Reference.objects.get(name="Referência fictícia Beta")
        self.assertEqual(created.current_stock, 0)
        self.assertIsNone(created.minimum)
        self.assertTrue(created.active)

    def test_editing_an_existing_historical_match_is_not_a_creation_or_merge(self):
        other = Reference.objects.create(
            code="SYN-HISTORICAL", name="REFERÊNCIA FICTÍCIA ALFA",
            location=self.other_location,
        )
        response = self.client.post(reverse("web:reference-edit", args=[other.pk]), self.data())
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Reference.objects.count(), 2)
        self.original.refresh_from_db()
        self.assertEqual(self.original.name, "  Referência   Fictícia Alfa  ")

    def test_matching_information_is_not_exposed_to_consultation_or_guests(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(self.create_url, {"check_name": self.original.name}).status_code, 403)
        self.assertEqual(self.client.post(self.create_url, self.data()).status_code, 403)
        self.assertEqual(Client().get(self.create_url, {"check_name": self.original.name}).status_code, 302)
        self.assertEqual(Reference.objects.count(), 1)


class HistoricalLocationImportTests(TestCase):
    def test_import_keeps_same_name_in_different_locations_separate(self):
        actor = get_user_model().objects.create_user(username="synthetic-history-admin", role="admin")
        with tempfile.TemporaryDirectory(prefix="synthetic-historical-locations-") as directory:
            path = synthetic_workbook(Path(directory) / "synthetic.xlsx")
            book = load_workbook(path)
            sheet = book["Stock"]
            sheet["A6"] = "CORREDOR DIREITO"
            for column in (1, 2, 3, 5, 7):
                sheet.cell(7, column).value = sheet.cell(5, column).value
            book.save(path)
            book.close()
            call_command("inspect_stock_excel", path, stage=True,
                         report=str(Path(directory) / "inspection.json"), stdout=io.StringIO())
        preview = build_preview(ImportBatch.objects.get().pk, {
            "schema_version": 1, "expected_total": 24, "expected_references": 2,
            "blank_opening_policy": "unknown_derived", "annotations": {}, "document_headers": {},
        }, actor=actor)
        report = publish_preview(preview.pk, actor=actor, confirmation_token=preview.confirmation_token)
        self.assertEqual(report["individual_mismatches"], 0)
        self.assertEqual(Reference.objects.count(), 2)
        self.assertEqual(Reference.objects.values("location").distinct().count(), 2)
        self.assertEqual(MovementLine.objects.count(), 2)
        self.assertEqual(len(matching_references("referência fictícia alfa")), 2)


@skipUnless(connection.vendor == "postgresql", "Creation serialization requires PostgreSQL")
class ConcurrentDuplicateCreationTests(TransactionTestCase):
    def test_simultaneous_creations_in_different_locations_cannot_both_create(self):
        actor = get_user_model().objects.create_user(username="synthetic-concurrent-match", role="admin")
        locations = [Location.objects.create(name=f"Local sintético {index}") for index in (1, 2)]
        barrier = Barrier(2)

        def create(index):
            connections.close_all()
            try:
                admin = get_user_model().objects.get(pk=actor.pk)
                location = Location.objects.get(pk=locations[index].pk)
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '8s'")
                    cursor.execute("SET statement_timeout = '10s'")
                barrier.wait(timeout=5)
                try:
                    save_reference(actor=admin, data={
                        "name": ["Referência Fictícia Concorrente", " referência  fictícia concorrente "][index],
                        "location": location,
                    })
                    return "created"
                except DuplicateReferenceError:
                    return "duplicate"
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(create, index) for index in (0, 1)]
            self.assertCountEqual([future.result(timeout=15) for future in futures], ["created", "duplicate"])
        self.assertEqual(Reference.objects.count(), 1)
        self.assertEqual(AuditEvent.objects.filter(action="reference.created").count(), 1)
