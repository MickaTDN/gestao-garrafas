"""Read-only inspection of the legacy Stock worksheet; never publishes balances.

Source values belong in private reports/staging, not in logs or test fixtures.
The numeric reconciliation is a candidate reproducing Excel SUM semantics.
Its assumptions and exceptions require review before any definitive import.
"""

import hashlib
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LOCATION_NAMES = {
    "corredor esquerdo": "Corredor Esquerdo",
    "corredor direito": "Corredor Direito",
    "patamar superior": "Patamar Superior Direito",
    "patamar superior direito": "Patamar Superior Direito",
    "patamar superior esquerdo": "Patamar Superior Esquerdo",
}


def normalized(value):
    text = " ".join(str(value or "").split()).casefold()
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


def integer(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    try:
        return int(value) if int(value) == value else None
    except (ValueError, OverflowError):
        return None


def serialized(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def cell_record(cell, cached_cell=None, coordinate=None):
    record = {
        "coordinate": getattr(cell, "coordinate", coordinate),
        "excel_type": cell.data_type,
        "python_type": type(cell.value).__name__,
        "value": serialized(cell.value),
    }
    if cell.data_type == "f":
        record["cached_value"] = serialized(cached_cell.value if cached_cell else None)
    return record


def calendar_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value.isoformat() if isinstance(value, date) else None


@dataclass
class InspectedRow:
    row_number: int
    raw_data: dict
    opening_stock: int | None
    reported_stock: int | None
    reconciled_stock: int | None


@dataclass
class Inspection:
    source_name: str
    sha256: str
    summary: dict = field(default_factory=dict)
    headers: dict = field(default_factory=dict)
    rows: list[InspectedRow] = field(default_factory=list)

    def as_report(self):
        return asdict(self)


def inspect_stock_workbook(path):
    """Inspect only Stock; preserve coordinates and source types without writing Excel."""
    path = Path(path)
    with path.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    inspection = Inspection(source_name=path.name, sha256=digest)
    issues = []

    def issue(code, cell=None, severity="warning"):
        issues.append({"code": code, "cell": cell, "severity": severity})

    formula_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached_book = None
    try:
        cached_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if "Stock" not in formula_book.sheetnames:
            issue("missing_stock_sheet", severity="error")
            inspection.summary = {"references": 0, "errors": 1, "warnings": 0, "issues": issues}
            return inspection
        sheet, cached_sheet = formula_book["Stock"], cached_book["Stock"]
        # Some Excel files retain inflated worksheet dimensions. Streaming the
        # declared range is safe; only populated header columns are interpreted.
        max_column = max(5, sheet.max_column or 5)
        formula_rows = sheet.iter_rows(max_col=max_column)
        cached_rows = cached_sheet.iter_rows(max_col=max_column)
        top = [next(formula_rows, ()) for _ in range(4)]
        cached_top = [next(cached_rows, ()) for _ in range(4)]
        if any(len(row) < 5 for row in top):
            issue("missing_stock_headers", severity="error")
            inspection.summary = {"references": 0, "errors": 1, "warnings": 0, "issues": issues}
            return inspection
        if normalized(top[1][4].value) != "inventario":
            issue("invalid_opening_header", "E2", "error")
        opening_date = calendar_date(top[2][4].value)
        if opening_date is None:
            issue("invalid_opening_date", "E3", "error")
        movement_columns = [
            j for j in range(6, max_column) if any(top[i][j].value is not None for i in (1, 2, 3))
        ]
        headers = {}
        dates = []
        for j in movement_columns:
            letter = get_column_letter(j + 1)
            when = calendar_date(top[2][j].value)
            if when is None:
                issue("invalid_movement_date", f"{letter}3", "error")
            else:
                dates.append(when)
            if not isinstance(top[1][j].value, str) or not top[1][j].value.strip():
                issue("invalid_movement_label", f"{letter}2", "error")
            if not isinstance(top[3][j].value, str) or not top[3][j].value.strip():
                issue("invalid_person_header", f"{letter}4", "error")
            headers[letter] = {
                "column": j + 1,
                "label": cell_record(top[1][j], cached_top[1][j], f"{letter}2"),
                "date": cell_record(top[2][j], cached_top[2][j], f"{letter}3"),
                "person": cell_record(top[3][j], cached_top[3][j], f"{letter}4"),
                "calendar_date": when,
            }
        if any(a < b for a, b in zip(dates, dates[1:])):
            issue("columns_not_reverse_chronological")
        inspection.headers = {
            "opening_label": cell_record(top[1][4], cached_top[1][4]),
            "opening_date": cell_record(top[2][4], cached_top[2][4]),
            "opening_calendar_date": opening_date,
            "movement_columns": headers,
        }
        location = None
        source_location = None
        initial_location = top[3][0].value
        if normalized(initial_location) in LOCATION_NAMES:
            location = LOCATION_NAMES[normalized(initial_location)]
            source_location = cell_record(top[3][0])
        seen = {}
        counts = Counter()
        sign_counts = Counter()
        for row_number, (row, cached_row) in enumerate(zip(formula_rows, cached_rows), start=5):
            name = row[0].value
            if normalized(name) in LOCATION_NAMES:
                location = LOCATION_NAMES[normalized(name)]
                source_location = cell_record(row[0])
                continue
            if name is None:
                if any(row[j].value is not None for j in (1, 2, 4, *movement_columns)):
                    issue("data_without_reference", f"A{row_number}", "error")
                continue
            if not isinstance(name, str) or not name.strip():
                issue("invalid_reference_name", f"A{row_number}", "error")
            if location is None:
                issue("reference_without_location", f"A{row_number}", "error")
            key = (location, normalized(name))
            if key in seen:
                issue("duplicate_reference_review_required", f"A{row_number}")
                counts["duplicates"] += 1
            seen[key] = row_number
            capacity = integer(row[1].value)
            if row[1].value is None:
                counts["missing_capacities"] += 1
            elif capacity is None or capacity <= 0:
                issue("invalid_box_capacity", f"B{row_number}", "error")
            opening = integer(row[4].value)
            if row[4].value is None:
                issue("blank_opening_assumed_zero_for_candidate", f"E{row_number}")
                counts["blank_openings"] += 1
            elif opening is None:
                issue("invalid_opening_quantity", f"E{row_number}", "error")
            reported = integer(cached_row[2].value)
            if reported is None:
                issue("missing_or_invalid_cached_stock", f"C{row_number}", "error")
            candidate = opening or 0
            entries = []
            if len(row) > 5 and row[5].value is not None:
                issue("unexpected_separator_value", f"F{row_number}", "error")
            row_movement_columns = sorted(
                set(movement_columns)
                | {j for j in range(6, max_column) if row[j].value is not None}
            )
            for j in row_movement_columns:
                cell = row[j]
                if cell.value is None:
                    continue
                letter = get_column_letter(j + 1)
                if letter not in headers:
                    # Preserve orphaned data instead of silently dropping a
                    # column whose three metadata cells have been cleared.
                    issue("movement_without_headers", f"{letter}2", "error")
                    movement_columns.append(j)
                    headers[letter] = {
                        "column": j + 1,
                        "label": cell_record(top[1][j], cached_top[1][j], f"{letter}2"),
                        "date": cell_record(top[2][j], cached_top[2][j], f"{letter}3"),
                        "person": cell_record(top[3][j], cached_top[3][j], f"{letter}4"),
                        "calendar_date": None,
                    }
                value = integer(cell.value)
                if value is None:
                    issue("nonnumeric_movement", cell.coordinate, "error")
                    counts["nonnumeric_movements"] += 1
                else:
                    candidate += value
                    label = normalized(top[1][j].value)
                    sign_counts[
                        f"{label}:{'positive' if value > 0 else 'negative' if value < 0 else 'zero'}"
                    ] += 1
                entries.append({"cell": cell_record(cell, cached_row[j]), "header_column": letter})
            if reported is not None and candidate != reported:
                issue("stock_reconciliation_mismatch", f"C{row_number}", "error")
                counts["mismatches"] += 1
            counts["reported_total"] += reported or 0
            counts["candidate_total"] += candidate
            counts["movement_cells"] += len(entries)
            counts[f"location:{location}"] += 1
            inspection.rows.append(
                InspectedRow(
                    row_number=row_number,
                    opening_stock=opening,
                    reported_stock=reported,
                    reconciled_stock=candidate,
                    raw_data={
                        "sheet": "Stock",
                        "source_location": source_location,
                        "mapped_location": location,
                        "sector": None,
                        "cells": {
                            get_column_letter(j + 1): cell_record(
                                row[j], cached_row[j], f"{get_column_letter(j + 1)}{row_number}"
                            )
                            for j in range(min(6, len(row)))
                        },
                        "movements": entries,
                        "reconciliation": {
                            "candidate_only": True,
                            "rule": "opening + sum of literal integer movements; blanks/text contribute zero as Excel SUM",
                            "requires_review": True,
                        },
                    },
                )
            )
        inspection.summary = {
            "schema_version": 1,
            "sheet": "Stock",
            "references": len(inspection.rows),
            "unique_references": len(seen),
            "duplicate_rows": counts["duplicates"],
            "missing_capacities": counts["missing_capacities"],
            "blank_openings": counts["blank_openings"],
            "movement_columns": len(movement_columns),
            "movement_cells": counts["movement_cells"],
            "nonnumeric_movements": counts["nonnumeric_movements"],
            "reported_stock_total": counts["reported_total"],
            "candidate_stock_total": counts["candidate_total"],
            "reconciliation_mismatches": counts["mismatches"],
            "earliest_movement_date": min(dates) if dates else None,
            "latest_movement_date": max(dates) if dates else None,
            "duplicate_dates": len(dates) - len(set(dates)),
            "signed_cells_by_source_label": dict(sign_counts),
            "candidate_only": True,
            "ready_for_publication": False,
            "errors": sum(i["severity"] == "error" for i in issues),
            "warnings": sum(i["severity"] == "warning" for i in issues),
            "issues": issues,
        }
        return inspection
    finally:
        formula_book.close()
        if cached_book is not None:
            cached_book.close()
