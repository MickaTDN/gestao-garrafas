"""Only generated workbooks and invented identities exercise definitive imports."""

import copy
import hashlib
import io
import tempfile
from datetime import datetime
from pathlib import Path
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError, connection, transaction
from django.test import TestCase
from openpyxl import load_workbook

from imports.reconciliation import build_preview, publish_preview
from imports.tests.test_excel import synthetic_workbook
from inventory.models import (
    AuditEvent,
    ImportBatch,
    ImportPreview,
    LegacyRow,
    Location,
    Movement,
    MovementLine,
    MovementLineSource,
    OpeningBalance,
    Person,
    Reference,
    Sector,
)


class ReconciledImportTests(TestCase):
    def setUp(self):
        self.private = tempfile.TemporaryDirectory(prefix="synthetic-reconciliation-")
        self.addCleanup(self.private.cleanup)
        self.counter = 0
        users = get_user_model()
        self.actor = users.objects.create_user(username="synthetic-import-admin", role="admin")
        self.viewer = users.objects.create_user(username="synthetic-import-viewer", role="viewer")

    def stage(self, update=None, **kwargs):
        self.counter += 1
        path = Path(self.private.name) / f"synthetic-{self.counter}.xlsx"
        synthetic_workbook(path, **kwargs)
        if update is not None:
            workbook = load_workbook(path)
            update(workbook["Stock"])
            workbook.save(path)
            workbook.close()
        try:
            call_command("inspect_stock_excel", path, stage=True,
                         report=str(path.with_suffix(".json")), stdout=io.StringIO())
        except CommandError:
            # Known source exceptions intentionally leave an unpublished batch.
            pass
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        return ImportBatch.objects.get(sha256=digest)

    def decisions(self, **overrides):
        result = {
            "schema_version": 1, "expected_total": 12, "expected_references": 1,
            "blank_opening_policy": "unknown_derived",
            "annotations": {}, "document_headers": {},
        }
        result.update(overrides)
        return result

    def preview(self, batch=None, decisions=None):
        return build_preview((batch or self.stage()).pk,
                             decisions or self.decisions(), actor=self.actor)

    def publish(self, preview):
        return publish_preview(preview.pk, actor=self.actor,
                               confirmation_token=preview.confirmation_token)

    def assert_unpublished(self):
        self.assertFalse(Reference.objects.exists())
        self.assertFalse(Movement.objects.exists())
        self.assertFalse(MovementLine.objects.exists())
        self.assertFalse(OpeningBalance.objects.exists())

    def test_preview_is_persistent_and_cannot_publish_implicitly(self):
        preview = self.preview()
        self.assertIsInstance(preview, ImportPreview)
        self.assertTrue(preview.summary["ready"])
        self.assertEqual(preview.summary["total_stock"], 12)
        self.assertEqual(preview.summary["reference_count"], 1)
        self.assertEqual(preview.rows[0]["stock"], 12)
        self.assertEqual(preview.import_batch.status, ImportBatch.Status.RECONCILED)
        self.assert_unpublished()
        self.assertEqual(AuditEvent.objects.get().action, "import.previewed")

    def test_publication_imports_ledger_and_official_stock_exactly_once(self):
        preview = self.preview()
        report = self.publish(preview)
        reference = Reference.objects.get()
        self.assertEqual(reference.current_stock, 12)  # Not 12 + the historical +2.
        self.assertEqual((reference.box_capacity, reference.minimum, reference.sector_id), (6, None, None))
        self.assertEqual(OpeningBalance.objects.get().quantity, 10)
        self.assertEqual(OpeningBalance.objects.get().reconciliation_quantity, 10)
        movement, line = Movement.objects.get(), MovementLine.objects.get()
        self.assertEqual((movement.kind, movement.status), ("legacy", "confirmed"))
        self.assertEqual(movement.date.isoformat(), "2026-01-07")
        self.assertIsNone(movement.occurred_at)
        self.assertEqual(line.quantity, 2)
        self.assertIsNone(line.stock_before)
        self.assertIsNone(line.stock_after)
        self.assertEqual(report["individual_mismatches"], 0)
        self.assertEqual(report["validation_rows"][0]["stock"], 12)
        self.assertTrue(report["published"])
        self.assertFalse(report["already_published"])
        preview.refresh_from_db()
        self.assertIsNotNone(preview.published_at)
        self.assertEqual(ImportBatch.objects.get().status, ImportBatch.Status.IMPORTED)

    def test_repeated_publication_is_a_noop_and_records_no_duplicate_ledger(self):
        preview = self.preview()
        first = self.publish(preview)
        counts = (Reference.objects.count(), Movement.objects.count(),
                  MovementLine.objects.count(), OpeningBalance.objects.count(),
                  MovementLineSource.objects.count(), AuditEvent.objects.count())
        again = self.publish(preview)
        self.assertTrue(again["already_published"])
        self.assertEqual(first["reconciliation_digest"], again["reconciliation_digest"])
        self.assertEqual(counts, (Reference.objects.count(), Movement.objects.count(),
                                 MovementLine.objects.count(), OpeningBalance.objects.count(),
                                 MovementLineSource.objects.count(), AuditEvent.objects.count()))
        self.assertEqual(Reference.objects.get().current_stock, 12)

    def test_wrong_confirmation_token_blocks_all_writes(self):
        preview = self.preview()
        with self.assertRaises(ValidationError):
            publish_preview(preview.pk, actor=self.actor, confirmation_token="different-token")
        self.assert_unpublished()

    def test_consultation_profile_cannot_prepare_or_publish(self):
        batch = self.stage()
        with self.assertRaises(PermissionDenied):
            build_preview(batch.pk, self.decisions(), actor=self.viewer)
        preview = self.preview(batch)
        with self.assertRaises(PermissionDenied):
            publish_preview(preview.pk, actor=self.viewer,
                            confirmation_token=preview.confirmation_token)
        self.assert_unpublished()

    def test_inactive_administrator_cannot_publish(self):
        preview = self.preview()
        self.actor.is_active = False
        with self.assertRaises(PermissionDenied):
            self.publish(preview)
        self.assert_unpublished()

    def test_blank_opening_remains_unknown_with_separate_reconciliation_balance(self):
        batch = self.stage(opening=None, stock=5)
        preview = self.preview(batch, self.decisions(expected_total=5))
        self.assertTrue(preview.summary["ready"])
        self.assertEqual(preview.summary["unknown_openings"], 1)
        self.assertFalse(preview.rows[0]["opening_known"])
        self.assertIsNone(preview.rows[0]["opening_quantity"])
        self.assertEqual(preview.rows[0]["reconciliation_quantity"], 3)
        self.publish(preview)
        opening = OpeningBalance.objects.get()
        self.assertIsNone(opening.quantity)
        self.assertEqual(opening.reconciliation_quantity, 3)
        self.assertEqual(opening.basis, "derived_unknown")
        self.assertIn("não representa uma contagem confirmada", opening.explanation)
        self.assertIsNone(LegacyRow.objects.get().raw_data["cells"]["E"]["value"])
        self.assertEqual(Reference.objects.get().current_stock, 5)

    def test_blank_openings_require_explicit_policy(self):
        batch = self.stage(opening=None, stock=2)
        preview = self.preview(batch, self.decisions(expected_total=2, blank_opening_policy=None))
        self.assertFalse(preview.summary["ready"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_annotation_has_no_quantitative_impact_and_preserves_original_text(self):
        batch = self.stage(lambda sheet: sheet.__setitem__("G5", "NOTA FICTÍCIA"), stock=10)
        resolutions = self.decisions(expected_total=10, annotations={
            "G5": {"original_value": "NOTA FICTÍCIA", "reason": "Anotação confirmada sem quantidade."},
        })
        preview = self.preview(batch, resolutions)
        self.assertTrue(preview.summary["ready"])
        self.assertEqual(preview.summary["annotations"], 1)
        self.publish(preview)
        source = MovementLineSource.objects.get()
        self.assertIsNone(source.quantity)
        self.assertEqual(source.raw_value, "NOTA FICTÍCIA")
        self.assertEqual(source.coordinate, "G5")
        self.assertEqual(MovementLine.objects.get().quantity, 0)
        self.assertEqual(Reference.objects.get().current_stock, 10)
        note = Movement.objects.get().source_metadata["annotations"][0]
        self.assertEqual(note["value"], "NOTA FICTÍCIA")
        self.assertEqual(note["decision"], resolutions["annotations"]["G5"]["reason"])

    def test_annotation_decision_must_match_source_exactly(self):
        batch = self.stage(lambda sheet: sheet.__setitem__("G5", "NOTA FICTÍCIA"), stock=10)
        preview = self.preview(batch, self.decisions(expected_total=10, annotations={
            "G5": {"original_value": "outro texto", "reason": "nota"},
        }))
        self.assertFalse(preview.summary["ready"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_numeric_document_header_does_not_create_an_invented_person(self):
        batch = self.stage(lambda sheet: sheet.__setitem__("G4", 123))
        preview = self.preview(batch, self.decisions(document_headers={
            "G": {"original_value": 123, "document_number": "123", "reason": "Documento confirmado."},
        }))
        self.assertEqual(preview.summary["unidentified_people"], 1)
        self.publish(preview)
        movement = Movement.objects.get()
        self.assertIsNone(movement.person_id)
        self.assertEqual(movement.document_number, "123")
        self.assertEqual(movement.source_metadata["header"]["person"]["value"], 123)
        self.assertEqual(movement.source_metadata["header"]["person"]["python_type"], "int")
        self.assertFalse(Person.objects.exists())

    def test_date_formatted_document_preserves_header_type_and_business_date(self):
        batch = self.stage(lambda sheet: sheet.__setitem__("G4", datetime(2040, 3, 4)))
        original = batch.summary["source_headers"]["movement_columns"]["G"]["person"]["value"]
        preview = self.preview(batch, self.decisions(document_headers={
            "G": {"original_value": original, "document_number": "DOC-FICTICIO-9", "reason": "Identificação documental."},
        }))
        self.publish(preview)
        movement = Movement.objects.get()
        self.assertEqual(movement.date.isoformat(), "2026-01-07")
        self.assertIsNone(movement.person_id)
        self.assertEqual(movement.document_number, "DOC-FICTICIO-9")
        self.assertEqual(movement.source_metadata["header"]["person"]["python_type"], "datetime")

    def test_duplicate_references_merge_without_losing_each_source_contribution(self):
        def duplicate(sheet):
            for column in (1, 2, 3, 5, 7):
                sheet.cell(6, column).value = sheet.cell(5, column).value
            sheet["A6"] = "  referência   fictícia alfa "

        batch = self.stage(duplicate)
        preview = self.preview(batch, self.decisions(expected_total=24))
        self.assertTrue(preview.summary["ready"])
        self.assertEqual(preview.summary["reference_count"], 1)
        self.assertEqual(preview.summary["row_count"], 2)
        self.assertEqual(preview.rows[0]["source_rows"], [5, 6])
        self.publish(preview)
        self.assertEqual(Reference.objects.count(), 1)
        self.assertEqual(Reference.objects.get().current_stock, 24)
        self.assertEqual(OpeningBalance.objects.count(), 2)
        self.assertEqual(MovementLine.objects.get().quantity, 4)
        sources = MovementLineSource.objects.order_by("coordinate")
        self.assertEqual(list(sources.values_list("coordinate", "quantity")), [("G5", 2), ("G6", 2)])
        self.assertEqual(LegacyRow.objects.filter(reference=Reference.objects.get()).count(), 2)

    def test_duplicate_conflicting_box_capacities_block_publication(self):
        def duplicate(sheet):
            for column in (1, 2, 3, 5, 7):
                sheet.cell(6, column).value = sheet.cell(5, column).value
            sheet["B6"] = 3

        preview = self.preview(self.stage(duplicate), self.decisions(expected_total=24))
        self.assertFalse(preview.summary["ready"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_unknown_box_capacity_is_preserved_without_assigning_a_sector(self):
        batch = self.stage(lambda sheet: sheet.__setitem__("B5", None))
        preview = self.preview(batch)
        self.publish(preview)
        reference = Reference.objects.get()
        self.assertIsNone(reference.box_capacity)
        self.assertIsNone(reference.sector_id)
        self.assertEqual(Location.objects.count(), 4)
        self.assertEqual(Sector.objects.count(), 27)

    def test_inverted_historical_signs_are_never_rewritten(self):
        def corrections(sheet):
            sheet["G5"] = -2
            sheet["H2"] = "Saida"
            sheet["H3"] = datetime(2026, 1, 7)
            sheet["H4"] = "Outra pessoa fictícia"
            sheet["H5"] = 4

        preview = self.preview(self.stage(corrections))
        self.assertEqual(preview.summary["negative_entries"], 1)
        self.assertEqual(preview.summary["positive_exits"], 1)
        self.publish(preview)
        quantities = dict(MovementLine.objects.values_list("movement__original_label", "quantity"))
        self.assertEqual(quantities, {"Entrada": -2, "Saida": 4})
        self.assertEqual(Reference.objects.get().current_stock, 12)

    def test_raw_source_change_invalidates_previous_confirmation(self):
        preview = self.preview()
        row = LegacyRow.objects.get()
        raw = copy.deepcopy(row.raw_data)
        raw["cells"]["A"]["value"] = "DESIGNAÇÃO SINTÉTICA ALTERADA"
        row.raw_data = raw
        row.save(update_fields=["raw_data"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_target_location_change_invalidates_previous_confirmation(self):
        preview = self.preview()
        location = Location.objects.get(name="Corredor Esquerdo")
        location.active = False
        location.save(update_fields=["active"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_stored_preview_rows_cannot_be_tampered_after_review(self):
        preview = self.preview()
        rows = copy.deepcopy(preview.rows)
        rows[0]["stock"] = 999
        ImportPreview.objects.filter(pk=preview.pk).update(rows=rows)
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_stored_preview_summary_cannot_be_tampered_after_review(self):
        preview = self.preview()
        summary = {**preview.summary, "total_stock": 999}
        ImportPreview.objects.filter(pk=preview.pk).update(summary=summary)
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_decisions_cannot_be_changed_without_a_new_preview(self):
        preview = self.preview()
        ImportPreview.objects.filter(pk=preview.pk).update(decisions=self.decisions(expected_total=999))
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_invalid_official_total_produces_blocked_preview_not_a_partial_import(self):
        preview = self.preview(decisions=self.decisions(expected_total=99))
        self.assertFalse(preview.summary["ready"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assert_unpublished()

    def test_publication_failure_rolls_back_the_complete_catalogue_and_ledger(self):
        preview = self.preview()
        with patch("imports.reconciliation.MovementLineSource.objects.create", side_effect=ValidationError("Falha sintética")):
            with self.assertRaises(ValidationError):
                self.publish(preview)
        self.assert_unpublished()
        self.assertFalse(Person.objects.exists())
        preview.refresh_from_db()
        self.assertIsNone(preview.published_at)
        self.assertEqual(ImportBatch.objects.get().status, ImportBatch.Status.RECONCILED)
        self.assertIsNone(LegacyRow.objects.get().reference_id)

    def test_another_file_cannot_add_a_second_snapshot_to_existing_stock(self):
        self.publish(self.preview())
        batch = self.stage(lambda sheet: sheet.__setitem__("A5", "OUTRA REFERÊNCIA FICTÍCIA"))
        preview = self.preview(batch)
        self.assertFalse(preview.summary["ready"])
        with self.assertRaises(ValidationError):
            self.publish(preview)
        self.assertEqual(Reference.objects.count(), 1)
        self.assertEqual(Reference.objects.get().current_stock, 12)
        self.assertEqual(Movement.objects.count(), 1)

    @skipUnless(connection.vendor == "postgresql", "Source SQL guards require PostgreSQL")
    def test_confirmed_source_provenance_cannot_be_rewritten_with_raw_sql(self):
        self.publish(self.preview())
        source = MovementLineSource.objects.get()
        with self.assertRaises(IntegrityError), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute("UPDATE inventory_movementlinesource SET quantity=%s WHERE id=%s", [99, source.pk])
        with self.assertRaises(IntegrityError), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute("DELETE FROM inventory_movementlinesource WHERE id=%s", [source.pk])
        with self.assertRaises(IntegrityError), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO inventory_movementlinesource (line_id, legacy_row_id, coordinate, quantity, raw_value) VALUES (%s,%s,%s,%s,%s)",
                [source.line_id, source.legacy_row_id, "H5", 1, "1"],
            )
        source.refresh_from_db()
        self.assertEqual(source.quantity, 2)
