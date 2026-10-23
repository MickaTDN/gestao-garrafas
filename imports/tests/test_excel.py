import io
import json
import os
import tempfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from django.conf import settings
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook

from imports.excel import inspect_stock_workbook
from inventory.models import ImportBatch, LegacyRow, Movement, Reference


def synthetic_workbook(path, *, opening=10, stock=12, formula=False, cache=None):
    """Invented fixture generated in a temporary directory; never real stock data."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Stock"
    sheet["A4"] = "CORREDOR ESQUERDO"
    sheet["B4"] = "UN. CX."
    sheet["C4"] = "UN. Artigo"
    sheet["D4"] = "Cx. Artigo"
    sheet["E2"] = "Inventário"
    sheet["E3"] = datetime(2025, 1, 6)
    sheet["G2"] = "Entrada"
    sheet["G3"] = datetime(2026, 1, 7)
    sheet["G4"] = "Pessoa fictícia"
    sheet["A5"] = "REFERÊNCIA FICTÍCIA ALFA"
    sheet["B5"] = 6
    sheet["C5"] = "=SUM(E5:G5)" if formula else stock
    sheet["E5"] = opening
    sheet["G5"] = 2
    book.save(path)
    if formula and cache is not None:
        # openpyxl never evaluates formulae. Supply a synthetic cached result
        # through XML so both formula-preservation and stale-cache detection run.
        with ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        tree = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
        cell = tree.find("s:sheetData/s:row/s:c[@r='C5']", ns)
        cell.find("s:v", ns).text = str(cache)
        entries["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            for name, contents in entries.items():
                archive.writestr(name, contents)
    return path


class ExcelInspectionTests(SimpleTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="synthetic-stock-")
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "synthetic.xlsx"

    def write_book(self, update=None, **kwargs):
        synthetic_workbook(self.path, **kwargs)
        if update:
            from openpyxl import load_workbook

            book = load_workbook(self.path)
            update(book["Stock"])
            book.save(self.path)

    def codes(self, inspection):
        return {i["code"] for i in inspection.summary["issues"]}

    def test_read_only_reconciliation_and_formula_cache_provenance(self):
        self.write_book(formula=True, cache=12)
        before = self.path.read_bytes()
        result = inspect_stock_workbook(self.path)
        self.assertEqual(result.summary["errors"], 0)
        self.assertEqual(result.rows[0].reconciled_stock, 12)
        self.assertEqual(result.rows[0].reported_stock, 12)
        self.assertEqual(result.rows[0].raw_data["cells"]["C"]["value"], "=SUM(E5:G5)")
        self.assertEqual(result.rows[0].raw_data["cells"]["C"]["cached_value"], 12)
        self.assertEqual(result.headers["opening_calendar_date"], "2025-01-06")
        self.assertEqual(self.path.read_bytes(), before)
        self.assertFalse(result.summary["ready_for_publication"])

    def test_missing_formula_cache_is_an_error_not_assumed_stock(self):
        self.write_book(formula=True)
        result = inspect_stock_workbook(self.path)
        self.assertIsNone(result.rows[0].reported_stock)
        self.assertIn("missing_or_invalid_cached_stock", self.codes(result))

    def test_stale_formula_cache_is_reconciliation_error(self):
        self.write_book(formula=True, cache=99)
        result = inspect_stock_workbook(self.path)
        self.assertEqual(result.summary["reconciliation_mismatches"], 1)

    def test_preserves_source_signs_and_repeated_dates(self):
        def update(sheet):
            sheet["G5"] = -2
            sheet["H2"] = "Saida"
            sheet["H3"] = datetime(2026, 1, 7)
            sheet["H4"] = "Outra pessoa fictícia"
            sheet["H5"] = 4

        self.write_book(update)
        result = inspect_stock_workbook(self.path)
        self.assertEqual(result.summary["duplicate_dates"], 1)
        self.assertEqual(
            [m["cell"]["value"] for m in result.rows[0].raw_data["movements"]], [-2, 4]
        )
        self.assertEqual(result.headers["movement_columns"]["G"]["calendar_date"], "2026-01-07")
        self.assertEqual(result.summary["errors"], 0)

    def test_text_and_blank_opening_are_explicit_unpublished_assumptions(self):
        self.write_book(
            lambda sheet: sheet.__setitem__("G5", "NOTA FICTÍCIA"), opening=None, stock=0
        )
        result = inspect_stock_workbook(self.path)
        self.assertIn("nonnumeric_movement", self.codes(result))
        self.assertIn("blank_opening_assumed_zero_for_candidate", self.codes(result))
        self.assertIsNone(result.rows[0].opening_stock)
        self.assertEqual(result.rows[0].reconciled_stock, 0)
        self.assertTrue(result.rows[0].raw_data["reconciliation"]["candidate_only"])
        self.assertIsNone(result.rows[0].raw_data["cells"]["E"]["value"])
        self.assertEqual(result.rows[0].raw_data["cells"]["E"]["coordinate"], "E5")

    def test_invalid_person_header_and_dates_preserve_raw_types(self):
        def update(sheet):
            sheet["G4"] = 123
            sheet["G3"] = "data por rever"

        self.write_book(update)
        result = inspect_stock_workbook(self.path)
        self.assertIn("invalid_person_header", self.codes(result))
        self.assertIn("invalid_movement_date", self.codes(result))
        self.assertEqual(result.headers["movement_columns"]["G"]["person"]["python_type"], "int")
        self.assertIsNone(result.headers["movement_columns"]["G"]["calendar_date"])

    def test_duplicate_reference_is_not_merged_or_dropped(self):
        def update(sheet):
            for col in (1, 2, 3, 5, 7):
                sheet.cell(6, col).value = sheet.cell(5, col).value

        self.write_book(update)
        result = inspect_stock_workbook(self.path)
        self.assertEqual(len(result.rows), 2)
        self.assertEqual(result.summary["unique_references"], 1)
        self.assertEqual(result.summary["duplicate_rows"], 1)
        self.assertIn("duplicate_reference_review_required", self.codes(result))

    def test_dates_out_of_order_preserve_column_identity(self):
        def update(sheet):
            sheet["H2"] = "Entrada"
            sheet["H3"] = datetime(2026, 1, 8)
            sheet["H4"] = "Pessoa fictícia"

        self.write_book(update)
        result = inspect_stock_workbook(self.path)
        self.assertIn("columns_not_reverse_chronological", self.codes(result))
        self.assertEqual(list(result.headers["movement_columns"]), ["G", "H"])

    def test_missing_stock_sheet_is_reportable(self):
        book = Workbook()
        book.save(self.path)
        result = inspect_stock_workbook(self.path)
        self.assertIn("missing_stock_sheet", self.codes(result))
        self.assertEqual(result.summary["errors"], 1)

    def test_orphaned_movement_column_is_preserved_and_flagged(self):
        self.write_book(lambda sheet: sheet.__setitem__("H5", 3), stock=15)
        result = inspect_stock_workbook(self.path)
        self.assertIn("movement_without_headers", self.codes(result))
        self.assertEqual(result.rows[0].reconciled_stock, 15)
        self.assertEqual(result.rows[0].raw_data["movements"][-1]["cell"]["coordinate"], "H5")
        self.assertEqual(result.headers["movement_columns"]["H"]["date"]["coordinate"], "H3")


class ExcelCommandTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="synthetic-private-report-")
        self.addCleanup(self.temp.cleanup)
        self.path = synthetic_workbook(Path(self.temp.name) / "synthetic.xlsx")
        self.report = Path(self.temp.name) / "inspection.json"

    def test_default_is_dry_run_and_output_does_not_contain_source_values(self):
        stdout = io.StringIO()
        call_command("inspect_stock_excel", self.path, report=str(self.report), stdout=stdout)
        self.assertFalse(ImportBatch.objects.exists())
        self.assertFalse(Reference.objects.exists())
        self.assertNotIn("REFERÊNCIA FICTÍCIA", stdout.getvalue())
        self.assertNotIn("Pessoa fictícia", stdout.getvalue())
        self.assertEqual(json.loads(self.report.read_text())["summary"]["references"], 1)
        if os.name == "posix":
            self.assertEqual(self.report.stat().st_mode & 0o777, 0o600)

    def test_staging_is_idempotent_and_never_publishes_balances(self):
        call_command(
            "inspect_stock_excel",
            self.path,
            report=str(self.report),
            stage=True,
            stdout=io.StringIO(),
        )
        call_command(
            "inspect_stock_excel",
            self.path,
            report=str(self.report.with_name("second.json")),
            stage=True,
            stdout=io.StringIO(),
        )
        self.assertEqual(ImportBatch.objects.count(), 1)
        self.assertEqual(LegacyRow.objects.count(), 1)
        self.assertIsNone(LegacyRow.objects.get().reference)
        self.assertEqual(ImportBatch.objects.get().status, ImportBatch.Status.STAGED)
        self.assertFalse(Reference.objects.exists())
        self.assertFalse(Movement.objects.exists())

    def test_reconciliation_errors_still_create_private_report_and_staged_source(self):
        synthetic_workbook(self.path, stock=99)
        with self.assertRaises(CommandError):
            call_command(
                "inspect_stock_excel",
                self.path,
                report=str(self.report),
                stage=True,
                stdout=io.StringIO(),
            )
        self.assertTrue(self.report.exists())
        self.assertEqual(ImportBatch.objects.get().status, ImportBatch.Status.STAGED)
        self.assertEqual(ImportBatch.objects.get().summary["reconciliation_mismatches"], 1)
        self.assertFalse(Movement.objects.exists())

    def test_reports_cannot_be_written_in_repository_even_if_ignored(self):
        with self.assertRaises(CommandError):
            call_command(
                "inspect_stock_excel",
                self.path,
                report=str(settings.BASE_DIR / "private" / "report.json"),
            )

    def test_existing_report_is_never_overwritten(self):
        self.report.write_text("preserve", encoding="utf-8")
        with self.assertRaises(CommandError):
            call_command("inspect_stock_excel", self.path, report=str(self.report))
        self.assertEqual(self.report.read_text(), "preserve")
