"""Explicit browser smoke suite: python manage.py test tests.browser."""

import os
import secrets
import shutil
from concurrent.futures import ThreadPoolExecutor

from django.contrib.staticfiles.handlers import StaticFilesHandler
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.db import connections
from django.test import SimpleTestCase, override_settings
from django.test.testcases import LiveServerThread
from django.urls import reverse
from playwright.sync_api import sync_playwright

from accounts.models import User
from inventory.models import Location, Reference


class BrowserSmokeTests(StaticLiveServerTestCase):
    @staticmethod
    def database_call(operation):
        # Playwright's synchronous driver runs an event loop in this thread.
        # Keep ORM work in a separate synchronous thread, with its own connection.
        def run():
            try:
                return operation()
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=1) as executor:
            return executor.submit(run).result()

    def setUp(self):
        self.password = secrets.token_urlsafe(32)
        self.viewer = User.objects.create_user(username="synthetic-browser", password=self.password)
        location = Location.objects.create(name="Local de teste sintético")
        self.location_id = location.pk
        Reference.objects.create(
            code="SYN-001",
            name="Referência fictícia azul",
            location=location,
            current_stock=15,
            box_capacity=6,
        )
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(
            executable_path=os.getenv("CHROMIUM_PATH") or shutil.which("chromium"),
            headless=True,
            args=["--no-sandbox"],
        )
        self.addCleanup(self.playwright.stop)
        self.addCleanup(self.browser.close)

    def test_desktop_and_mobile_authenticated_stock(self):
        self.database_call(
            lambda: Reference.objects.create(
                code="SYN-STOCK-OTHER",
                name="Outra garrafa fictícia",
                location_id=self.location_id,
                current_stock=1,
            )
        )
        for width, height in [(1366, 900), (390, 844)]:
            with self.subTest(width=width):
                page = self.browser.new_page(viewport={"width": width, "height": height})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                response = page.goto(self.live_server_url + reverse("web:stock"))
                self.assertEqual(response.status, 200)
                self.assertIn(reverse("login"), page.url)
                page.get_by_label("Utilizador", exact=True).fill(self.viewer.username)
                page.get_by_label("Palavra-passe", exact=True).fill(self.password)
                page.get_by_role("button", name="Entrar", exact=True).click()
                self.assertIn("Referência fictícia azul", page.locator("body").inner_text())
                # Submit through the real form, including query-string round trip.
                page.locator('input[name="q"]').fill("Referência fictícia azul")
                page.wait_for_function("() => document.querySelectorAll('tbody tr').length === 1")
                self.assertEqual(page.locator("tbody tr").count(), 1)
                self.assertIn("15", page.locator("tbody").inner_text())
                self.assertLessEqual(
                    page.evaluate("document.documentElement.scrollWidth"),
                    width,
                    "A página não deve transbordar horizontalmente em telemóvel.",
                )
                forbidden = page.goto(self.live_server_url + reverse("web:users"))
                self.assertEqual(forbidden.status, 403)
                self.assertEqual(errors, [])
                page.close()

    def test_administrator_creates_and_edits_reference_on_mobile(self):
        self.database_call(lambda: User.objects.filter(pk=self.viewer.pk).update(role="admin"))
        page = self.browser.new_page(viewport={"width": 390, "height": 844})
        page.goto(self.live_server_url + reverse("web:reference-create"))
        page.get_by_label("Utilizador", exact=True).fill(self.viewer.username)
        page.get_by_label("Palavra-passe", exact=True).fill(self.password)
        page.get_by_role("button", name="Entrar", exact=True).click()
        page.get_by_label("Designação", exact=True).fill("Referência fictícia verde")
        page.get_by_label("Garrafas por caixa", exact=True).fill("6")
        page.get_by_label("Localização", exact=True).select_option(str(self.location_id))
        self.assertFalse(
            page.get_by_label("Ativar alerta de stock mínimo", exact=True).is_checked()
        )
        page.get_by_role("button", name="Guardar", exact=True).click()
        self.assertIn("Referência fictícia verde", page.locator("h1").inner_text())
        reference = self.database_call(
            lambda: Reference.objects.get(name="Referência fictícia verde")
        )
        self.assertEqual(reference.current_stock, 0)
        self.assertIsNone(reference.minimum)
        page.goto(self.live_server_url + reverse("web:reference-edit", args=[reference.pk]))
        page.get_by_label("Ativar alerta de stock mínimo", exact=True).check()
        page.get_by_label("Stock mínimo em garrafas", exact=True).fill("4")
        page.get_by_role("button", name="Guardar", exact=True).click()
        self.database_call(reference.refresh_from_db)
        self.assertEqual(reference.minimum, 4)
        self.assertEqual(reference.current_stock, 0)
        self.assertLessEqual(page.evaluate("document.documentElement.scrollWidth"), 390)
        page.close()


class OperationalBrowserTests(StaticLiveServerTestCase):
    """Exercise actual forms and JS on both viewport sizes with synthetic fixtures."""

    database_call = staticmethod(BrowserSmokeTests.database_call)

    def setUp(self):
        self.password = secrets.token_urlsafe(32)
        self.admin = User.objects.create_user(
            username="synthetic-operations", password=self.password, role="admin"
        )
        self.location = Location.objects.create(name="Origem browser sintética")
        self.destination = Location.objects.create(name="Destino browser sintético")
        self.first = Reference.objects.create(
            code="SYN-OP-1",
            name="Azul sintética operacional",
            box_capacity=6,
            current_stock=15,
            location=self.location,
        )
        self.second = Reference.objects.create(
            code="SYN-OP-2",
            name="Verde sintética operacional",
            current_stock=5,
            location=self.location,
        )
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(
            executable_path=os.getenv("CHROMIUM_PATH") or shutil.which("chromium"),
            headless=True,
            args=["--no-sandbox"],
        )
        self.addCleanup(self.playwright.stop)
        self.addCleanup(self.browser.close)

    def page(self, kind, width):
        page = self.browser.new_page(
            viewport={"width": width, "height": 900 if width > 760 else 844}
        )
        self.addCleanup(page.close)
        page.set_default_timeout(6000)
        self.errors = []
        page.on("pageerror", lambda error: self.errors.append(str(error)))
        page.goto(self.live_server_url + reverse(f"web:{kind}"))
        page.get_by_label("Utilizador", exact=True).fill(self.admin.username)
        page.get_by_label("Palavra-passe", exact=True).fill(self.password)
        page.get_by_role("button", name="Entrar", exact=True).click()
        return page

    def assert_responsive(self, page, width):
        self.assertLessEqual(
            page.evaluate("document.documentElement.scrollWidth"),
            width,
            page.evaluate(
                "() => Array.from(document.querySelectorAll('body *')).filter(e => {const r=e.getBoundingClientRect(); return r.width && r.right > innerWidth;}).map(e => [e.tagName, e.className, Math.round(e.getBoundingClientRect().right)]).slice(0,15)"
            ),
        )
        self.assertEqual(self.errors, [])
        if directory := os.getenv("BROWSER_ARTIFACT_DIR"):
            os.makedirs(directory, exist_ok=True)
            self.screenshot_count = getattr(self, "screenshot_count", 0) + 1
            page.screenshot(
                path=os.path.join(directory, f"{self._testMethodName}-{self.screenshot_count}.png"),
                full_page=True,
            )

    def stocks(self):
        return self.database_call(
            lambda: list(Reference.objects.order_by("code").values_list("current_stock", flat=True))
        )

    def add_reference(self, page, name):
        page.get_by_label("Pesquisar referência", exact=True).fill(name)
        page.get_by_role("button", name=name, exact=True).click()

    def entry_exit(self, kind, width):
        page = self.page(kind, width)
        self.assertEqual(page.locator("#reference-results button").count(), 0)
        self.add_reference(page, "Azul sintética operacional")
        self.add_reference(page, "Verde sintética operacional")
        edit_url = page.url.split("?", 1)[0]
        self.assertEqual(
            page.locator(".operation-line").count(), 2, page.locator("body").inner_text()
        )
        first = page.locator(".operation-line").filter(has_text="Azul sintética operacional")
        second = page.locator(".operation-line").filter(has_text="Verde sintética operacional")
        for row in (first, second):
            self.assertEqual(row.get_by_label("Caixas", exact=True).input_value(), "")
            self.assertEqual(row.get_by_label("Garrafas soltas", exact=True).input_value(), "")
        page.get_by_label(
            "Pessoa que entregou" if kind == "entry" else "Pessoa que levantou", exact=True
        ).fill("Pessoa browser sintética")
        first.get_by_label("Caixas", exact=True).fill("1")
        first.get_by_label("Garrafas soltas", exact=True).fill("8" if kind == "entry" else "2")
        second.get_by_label("Garrafas soltas", exact=True).fill("3")
        self.assertEqual(
            first.locator(".line-total").inner_text(), "14" if kind == "entry" else "8"
        )
        self.assertIn("29" if kind == "entry" else "7", first.locator(".line-after").inner_text())
        self.assert_responsive(page, width)
        first.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertEqual(first.get_attribute("open"), None)
        second.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertEqual(second.get_attribute("open"), None)
        self.assertEqual(self.stocks(), [15, 5])
        page.get_by_role("button", name="Guardar operação", exact=True).click()
        self.assertEqual(self.stocks(), [15, 5])
        page.goto(self.live_server_url + reverse(f"web:{kind}"))
        page.locator(f'a[href="{edit_url.removeprefix(self.live_server_url)}"]').click()
        self.assertEqual(page.url, edit_url)
        self.assertEqual(page.locator('input[name="form-0-boxes"]').input_value(), "1")
        self.assertEqual(page.locator('input[name="form-1-boxes"]').input_value(), "")
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        self.assertIn("/conferir/", page.url)
        self.assertEqual(self.stocks(), [15, 5])
        # Correction before confirmation is available and invalidates the old review.
        page.get_by_role("link", name="Corrigir referências ou quantidades", exact=True).click()
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        if kind == "entry":
            page.locator('select[name="report_requested"]').select_option("yes")
        self.assert_responsive(page, width)
        page.locator('input[name="confirm"]').check()
        page.get_by_role(
            "button", name="Confirmar entrada" if kind == "entry" else "Confirmar saída", exact=True
        ).click()
        self.assertIn("CONFIRMADO", page.locator("body").inner_text())
        self.assertEqual(self.stocks(), [29, 8] if kind == "entry" else [7, 2])
        self.assertIn("Descarregar PDF guardado", page.locator("body").inner_text())
        self.assert_responsive(page, width)

    def inventory_flow(self, width):
        page = self.page("inventory", width)
        page.get_by_label("Âmbito", exact=True).select_option("full")
        page.get_by_role("button", name="Preparar inventário completo", exact=True).click()
        first = page.locator(".operation-line").filter(has_text="Azul sintética operacional")
        second = page.locator(".operation-line").filter(has_text="Verde sintética operacional")
        first.locator("summary").click()
        second.locator("summary").click()
        self.assertEqual(first.get_by_label("Caixas", exact=True).input_value(), "")
        self.assertEqual(second.get_by_label("Garrafas soltas", exact=True).input_value(), "")
        first.get_by_label("Caixas", exact=True).fill("2")
        first.get_by_label("Garrafas soltas", exact=True).fill("1")
        first.get_by_label("Motivo da diferença / transferência", exact=True).fill(
            "Contagem física browser sintética"
        )
        first.get_by_label("Localização após inventário", exact=True).select_option(
            str(self.destination.pk)
        )
        self.assertEqual(first.locator(".line-total").inner_text(), "13")
        self.assertEqual(first.locator(".line-difference").inner_text(), "-2")
        second.get_by_label("Garrafas soltas", exact=True).fill("0")
        second.get_by_label("Motivo da diferença / transferência", exact=True).fill(
            "Zero contado explicitamente"
        )
        self.assertEqual(second.locator(".line-total").inner_text(), "0")
        self.assert_responsive(page, width)
        page.get_by_role("button", name="Guardar operação", exact=True).click()
        self.assertEqual(self.stocks(), [15, 5])
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        self.assertIn("/conferir/", page.url)
        page.locator('input[name="confirm"]').check()
        page.get_by_role("button", name="Confirmar inventário", exact=True).click()
        self.assertIn("CONFIRMADO", page.locator("body").inner_text())
        self.assertEqual(self.stocks(), [13, 0])
        destination_id = self.database_call(
            lambda: Reference.objects.get(pk=self.first.pk).location_id
        )
        self.assertEqual(destination_id, self.destination.pk)
        self.assertIn("Transferência", page.locator("body").inner_text())
        self.assert_responsive(page, width)

    def test_entry_desktop(self):
        self.entry_exit("entry", 1366)

    def test_entry_mobile(self):
        self.entry_exit("entry", 390)

    def test_exit_desktop(self):
        self.entry_exit("exit", 1366)

    def test_exit_mobile(self):
        self.entry_exit("exit", 390)

    def a2_edit_history_and_pdf(self, width):
        from inventory.documents import verified_bytes
        from inventory.models import AuditEvent
        from tests.document_fixtures import confirmed_operation

        page = self.page("entry", width)
        self.add_reference(page, self.first.name)
        row = page.locator(".operation-line")
        row.get_by_label("Garrafas soltas", exact=True).fill("2")
        row.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertIsNone(row.get_attribute("open"))
        self.assertIn("2 garrafas", row.locator("summary").inner_text())
        row.locator("summary").click()
        row.get_by_label("Garrafas soltas", exact=True).fill("3")
        row.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertIn("3 garrafas", row.locator("summary").inner_text())
        row.locator("summary").click()
        row.get_by_role("button", name="Remover referência", exact=True).click()
        self.assertEqual(page.locator(".operation-line").count(), 0)
        self.assertEqual(self.stocks(), [15, 5])

        movement = self.database_call(
            lambda: confirmed_operation(self.admin, self.first, signed=True)
        )
        document = self.database_call(
            lambda: movement.documents.select_related("annual_number").get(kind="pdf")
        )
        original = self.database_call(lambda: verified_bytes(document))
        before = self.database_call(lambda: AuditEvent.objects.count())
        page.goto(self.live_server_url + reverse("web:history"))
        details = page.locator(".movement-details")
        self.assertIsNone(details.get_attribute("open"))
        self.assertFalse(details.locator("table").is_visible())
        details.locator("summary").click()
        self.assertTrue(details.locator("table").is_visible())
        details.locator("summary").click()
        self.assertFalse(details.locator("table").is_visible())
        self.assertEqual(self.database_call(lambda: AuditEvent.objects.count()), before)
        self.assert_responsive(page, width)

        page.goto(self.live_server_url + reverse("web:document-delete", args=[document.pk]))
        page.get_by_label("Confirmo a eliminação apenas deste ficheiro PDF.", exact=True).check()
        page.get_by_role("button", name="Eliminar PDF", exact=True).click()
        page.goto(self.live_server_url + reverse("web:history"))
        details = page.locator(".movement-details")
        details.locator("summary").click()
        self.assertIn("Ficheiro PDF eliminado", details.inner_text())
        details.get_by_role("button", name=f"Regenerar PDF {document.number}", exact=True).click()
        self.assertIn("Reprodução verificada byte a byte", page.locator("body").inner_text())
        self.assertEqual(self.database_call(lambda: verified_bytes(document)), original)
        self.assertEqual(self.stocks(), [13, 5])
        self.assert_responsive(page, width)

    def test_a2_edit_history_and_pdf_desktop(self):
        self.a2_edit_history_and_pdf(1366)

    def test_a2_edit_history_and_pdf_mobile(self):
        self.a2_edit_history_and_pdf(390)

    def test_inventory_desktop(self):
        self.inventory_flow(1366)

    def test_inventory_mobile(self):
        self.inventory_flow(390)

    def sector_and_new_reference_flow(self, width):
        from inventory.models import Sector

        origin_sector, destination_sector = self.database_call(
            lambda: (
                Sector.objects.create(location=self.location, name="Origem setor fictício"),
                Sector.objects.create(location=self.destination, name="Destino setor fictício"),
            )
        )
        page = self.page("entry", width)
        self.assertEqual(page.locator('select[name="location"]').count(), 0)
        self.add_reference(page, "Azul sintética operacional")
        edit_url = page.url
        page.get_by_label("Garrafas soltas", exact=True).fill("3")
        page.get_by_role("button", name="Criar nova referência", exact=True).click()
        self.assertIn("Nova referência na entrada", page.locator("h1").inner_text())
        page.get_by_label("Designação", exact=True).fill("Nova referência browser fictícia")
        location = page.get_by_label("Localização", exact=True)
        sector = page.get_by_label("Setor", exact=True)
        location.select_option(str(self.location.pk))
        sector.select_option(str(origin_sector.pk))
        location.select_option(str(self.destination.pk))
        self.assertEqual(sector.input_value(), "")
        self.assertEqual(
            sector.locator("option:not([disabled])").evaluate_all(
                "options => options.map(o => o.value)"
            ),
            ["", str(destination_sector.pk)],
        )
        sector.select_option(str(destination_sector.pk))
        self.assert_responsive(page, width)
        page.get_by_role("button", name="Criar referência", exact=True).click()
        self.assertEqual(page.url, edit_url)
        self.assertEqual(page.get_by_label("Garrafas soltas", exact=True).input_value(), "3")
        self.assertEqual(self.stocks(), [0, 15, 5])
        # The same dynamic selector works when editing the catalogue and counting.
        page.goto(self.live_server_url + reverse("web:reference-edit", args=[self.first.pk]))
        page.get_by_label("Localização", exact=True).select_option(str(self.destination.pk))
        page.get_by_label("Setor", exact=True).select_option(str(destination_sector.pk))
        page.get_by_role("button", name="Guardar", exact=True).click()
        page.goto(self.live_server_url + reverse("web:inventory"))
        self.add_reference(page, "Azul sintética operacional")
        sector = page.get_by_label("Setor após inventário", exact=True)
        self.assertEqual(sector.input_value(), str(destination_sector.pk))
        page.get_by_label("Localização após inventário", exact=True).select_option(
            str(self.location.pk)
        )
        self.assertEqual(sector.input_value(), "")
        self.assertEqual(
            sector.locator("option:not([disabled])").evaluate_all(
                "options => options.map(o => o.value)"
            ),
            ["", str(origin_sector.pk)],
        )
        self.assertEqual(self.stocks(), [0, 15, 5])
        self.assert_responsive(page, width)

    def test_sector_filter_and_new_reference_desktop(self):
        self.sector_and_new_reference_flow(1366)

    def test_sector_filter_and_new_reference_mobile(self):
        self.sector_and_new_reference_flow(390)

    def a3_inventory_partial(self, width):
        from inventory.models import Movement, Sector

        def prepare():
            sectors = {
                name: Sector.objects.create(location=self.location, name="Setor " + name)
                for name in ("1", "2", "9", "10")
            }
            target = Sector.objects.create(location=self.destination, name="3")
            Reference.objects.filter(pk=self.first.pk).update(sector=sectors["10"])
            Reference.objects.filter(pk=self.second.pk).update(sector=sectors["2"])
            return sectors, target

        sectors, target = self.database_call(prepare)
        page = self.page("inventory", width)
        search = page.locator("#reference-search")
        search.get_by_label("Localização", exact=True).select_option(str(self.location.pk))
        sector = search.get_by_label("Setor", exact=True)
        self.assertEqual(
            sector.locator("option").evaluate_all("options => options.map(o => o.value)"),
            ["", *(str(sectors[name].pk) for name in ("1", "2", "9", "10"))],
        )
        sector.select_option(str(sectors["10"].pk))
        self.add_reference(page, self.first.name)
        first = page.locator(".operation-line").filter(has_text=self.first.name)
        self.assertEqual(first.get_by_label("Caixas", exact=True).input_value(), "")
        self.assertEqual(first.get_by_label("Garrafas soltas", exact=True).input_value(), "")
        first.get_by_label("Caixas", exact=True).fill("2")
        first.get_by_label("Garrafas soltas", exact=True).fill("3")
        first.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertIsNone(first.get_attribute("open"))
        self.assertIn("2 caixas + 3 garrafas", first.locator("summary").inner_text())
        page.locator("#reference-search").get_by_label("Setor", exact=True).select_option("")
        self.add_reference(page, self.second.name)
        second = page.locator(".operation-line").filter(has_text=self.second.name)
        second.get_by_label("Garrafas soltas", exact=True).fill("0")
        second.get_by_label("Motivo da diferença / transferência", exact=True).fill(
            "Zero físico fictício A3"
        )
        second.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertIsNone(second.get_attribute("open"))
        self.assertIn("0 garrafas", second.locator("summary").inner_text())
        first.locator("summary").click()
        first.get_by_label("Garrafas soltas", exact=True).fill("2")
        first.get_by_label("Motivo da diferença / transferência", exact=True).fill(
            "Transferência e contagem fictícias A3"
        )
        target_sector = first.get_by_label("Setor após inventário", exact=True)
        self.assertEqual(target_sector.input_value(), str(sectors["10"].pk))
        first.get_by_label("Localização após inventário", exact=True).select_option(
            str(self.destination.pk)
        )
        self.assertEqual(target_sector.input_value(), "")
        self.assertEqual(
            target_sector.locator("option").evaluate_all("options => options.map(o => o.value)"),
            ["", str(target.pk)],
        )
        target_sector.select_option(str(target.pk))
        first.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        second.locator("summary").click()
        second.get_by_role("button", name="Remover referência", exact=True).click()
        self.assertEqual(page.locator(".operation-line").count(), 1)
        self.assertEqual(self.stocks(), [15, 5])
        self.assertEqual(
            self.database_call(lambda: Reference.objects.get(pk=self.first.pk).location_id),
            self.location.pk,
        )
        self.assert_responsive(page, width)
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        page.get_by_label("Validação do inventário", exact=True).select_option("signed")
        canvas = page.locator("#signature-canvas")
        canvas.scroll_into_view_if_needed()
        bounds = canvas.bounding_box()
        points = [
            (bounds["x"] + bounds["width"] * x, bounds["y"] + bounds["height"] * y)
            for x, y in ((0.1, 0.7), (0.25, 0.2), (0.5, 0.8), (0.8, 0.3))
        ]
        if width < 760:
            session = page.context.new_cdp_session(page)
            session.send("Emulation.setTouchEmulationEnabled", {"enabled": True})
            session.send(
                "Input.dispatchTouchEvent",
                {"type": "touchStart", "touchPoints": [{"x": points[0][0], "y": points[0][1]}]},
            )
            for x, y in points[1:]:
                session.send(
                    "Input.dispatchTouchEvent",
                    {"type": "touchMove", "touchPoints": [{"x": x, "y": y}]},
                )
            session.send("Input.dispatchTouchEvent", {"type": "touchEnd", "touchPoints": []})
            session.detach()
        else:
            page.mouse.move(*points[0])
            page.mouse.down()
            for point in points[1:]:
                page.mouse.move(*point, steps=5)
            page.mouse.up()
        page.locator('input[name="confirm"]').check()
        page.get_by_role("button", name="Confirmar inventário", exact=True).click()
        self.assertEqual(self.stocks(), [14, 5])
        movement = self.database_call(lambda: Movement.objects.latest("pk"))
        official = self.database_call(lambda: str(movement.report_number))
        self.assertIn(official, page.locator("h1").inner_text())
        self.assertEqual(self.database_call(lambda: movement.documents.count()), 2)
        self.assertEqual(
            self.database_call(lambda: Reference.objects.get(pk=self.first.pk).location_id),
            self.destination.pk,
        )
        self.assert_responsive(page, width)

    def test_a3_inventory_partial_desktop(self):
        self.a3_inventory_partial(1366)

    def test_a3_inventory_partial_mobile(self):
        self.a3_inventory_partial(390)

    def a3_password_flow(self, width):
        page = self.page("entry", width)
        page.goto(self.live_server_url + reverse("web:settings"))
        page.get_by_role("link", name="A minha conta", exact=True).click()
        self.assert_responsive(page, width)
        replacement = secrets.token_urlsafe(32)
        page.get_by_label("Palavra-passe atual", exact=True).fill(self.password)
        page.get_by_label("Nova palavra-passe", exact=True).fill(replacement)
        page.get_by_label("Confirmar nova palavra-passe", exact=True).fill(replacement)
        page.get_by_role("button", name="Alterar palavra-passe", exact=True).click()
        self.assertIn("Palavra-passe alterada", page.locator("body").inner_text())
        self.assertNotIn(replacement, page.content())
        viewer = self.database_call(
            lambda: User.objects.create_user(username="a3-browser-reset-synthetic")
        )
        page.goto(self.live_server_url + reverse("web:user-password-reset", args=[viewer.pk]))
        self.assert_responsive(page, width)
        temporary = secrets.token_urlsafe(32)
        page.get_by_label("A sua palavra-passe de Administrador", exact=True).fill(replacement)
        page.get_by_label("Palavra-passe temporária", exact=True).fill(temporary)
        page.get_by_label("Confirmar palavra-passe temporária", exact=True).fill(temporary)
        page.locator('input[name="confirm"]').check()
        page.get_by_role("button", name="Repor palavra-passe", exact=True).click()
        self.assertNotIn(temporary, page.content())
        page.get_by_role("button", name="Sair", exact=True).click()
        page.goto(self.live_server_url + reverse("login"))
        page.get_by_label("Utilizador", exact=True).fill(viewer.username)
        page.get_by_label("Palavra-passe", exact=True).fill(temporary)
        page.get_by_role("button", name="Entrar", exact=True).click()
        self.assertIn("É obrigatório alterar", page.locator("body").inner_text())
        self.assert_responsive(page, width)
        page.goto(self.live_server_url + reverse("web:stock"))
        self.assertIn(reverse("web:my-account"), page.url)
        final = secrets.token_urlsafe(32)
        page.get_by_label("Palavra-passe atual", exact=True).fill(temporary)
        page.get_by_label("Nova palavra-passe", exact=True).fill(final)
        page.get_by_label("Confirmar nova palavra-passe", exact=True).fill(final)
        page.get_by_role("button", name="Alterar palavra-passe", exact=True).click()
        page.goto(self.live_server_url + reverse("web:stock"))
        self.assertIn("Consultar stock", page.locator("h1").inner_text())
        self.assert_responsive(page, width)

    def test_a3_password_flow_desktop(self):
        self.a3_password_flow(1366)

    def test_a3_password_flow_mobile(self):
        self.a3_password_flow(390)

    def a4_worklist(self, width, full):
        from inventory.models import AuditEvent, Movement, Sector

        def prepare():
            nine = Sector.objects.create(location=self.location, name="Setor 9")
            ten = Sector.objects.create(location=self.location, name="Setor 10")
            Reference.objects.filter(pk=self.first.pk).update(sector=nine)
            Reference.objects.filter(pk=self.second.pk).update(sector=ten)

        self.database_call(prepare)
        page = self.page("inventory", width)
        search = page.locator("#reference-search")
        search.get_by_label("Localização", exact=True).select_option(str(self.location.pk))
        page.wait_for_function(
            "() => !document.querySelector('#reference-results').hasAttribute('aria-busy')"
        )
        results = page.locator("#reference-results")
        self.assertIn("Setor 9", results.inner_text())
        self.assertIn("Setor 10", results.inner_text())
        self.assertLess(
            results.inner_text().index("Setor 9"), results.inner_text().index("Setor 10")
        )
        self.assertNotIn(self.location.name, results.inner_text())
        if full:
            page.get_by_label("Âmbito", exact=True).select_option("full")
            page.get_by_role("button", name="Preparar inventário completo", exact=True).click()
            self.assertEqual(page.locator("details.operation-line[open]").count(), 0)
            first = page.locator(".operation-line").filter(has_text=self.first.name)
            second = page.locator(".operation-line").filter(has_text=self.second.name)
            self.assertIn("Setor 9", first.locator("summary").inner_text())
            self.assertIn("Ainda não contado", first.locator("summary").inner_text())
            first.locator("summary").click()
        else:
            page.get_by_role("button", name=self.first.name, exact=True).click()
            first = page.locator(".operation-line").filter(has_text=self.first.name)
            self.assertIsNotNone(first.get_attribute("open"))
        self.assertIn("location=" + str(self.location.pk), page.url)
        self.assertEqual(first.get_by_label("Caixas", exact=True).input_value(), "")
        self.assertEqual(first.get_by_label("Garrafas soltas", exact=True).input_value(), "")
        first.get_by_label("Garrafas soltas", exact=True).fill("0")
        first.get_by_label("Motivo da diferença / transferência", exact=True).fill(
            "Zero físico A4 fictício"
        )
        first.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertIsNone(first.get_attribute("open"))
        summary = first.locator("summary").inner_text()
        self.assertIn("Setor 9", summary)
        self.assertIn("0 garrafas", summary)
        before = self.database_call(lambda: AuditEvent.objects.count())
        if full:
            first.locator("summary").click()
        else:
            page.get_by_role(
                "link", name="Abrir contagem de " + self.first.name, exact=True
            ).click()
        self.assertIsNotNone(first.get_attribute("open"))
        self.assertEqual(self.database_call(lambda: AuditEvent.objects.count()), before)
        first.get_by_label("Garrafas soltas", exact=True).fill("14")
        first.get_by_role("button", name="Confirmar quantidade", exact=True).click()
        self.assertEqual(self.stocks(), [15, 5])
        self.assert_responsive(page, width)
        if full:
            first.locator("summary").click()
            first.get_by_role("button", name="Limpar contagem", exact=True).click()
            self.assertIn("Ainda não contado", first.locator("summary").inner_text())
            movement = self.database_call(lambda: Movement.objects.latest("pk"))
            self.assertEqual(self.database_call(lambda: movement.lines.count()), 2)
            self.assertEqual(
                self.database_call(
                    lambda: movement.lines.get(reference_id=self.first.pk).counted_quantity
                ),
                None,
            )
            second = page.locator(".operation-line").filter(has_text=self.second.name)
            second.locator("summary").click()
            self.assertEqual(second.get_by_label("Garrafas soltas", exact=True).input_value(), "")
        else:
            # Live name search keeps the saved count; location/sector are shown
            # when returning to an unrestricted name lookup.
            page.locator("#reference-search").get_by_label("Localização", exact=True).select_option(
                ""
            )
            page.get_by_label("Pesquisar referência", exact=True).fill(self.first.name)
            link = page.get_by_role("link", name="Abrir contagem de " + self.first.name, exact=True)
            page.wait_for_function(
                "() => !document.querySelector('#reference-results').hasAttribute('aria-busy')"
            )
            self.assertIn(self.location.name, link.inner_text())
            self.assertIn("Setor 9", link.inner_text())
            self.assertIn("14 garrafas", link.inner_text())
            link.click()
            first.get_by_role("button", name="Remover referência", exact=True).click()
            self.assertEqual(page.locator(".operation-line").count(), 0)
        self.assertEqual(self.stocks(), [15, 5])
        self.assert_responsive(page, width)

    def test_a4_partial_worklist_desktop(self):
        self.a4_worklist(1366, False)

    def test_a4_partial_worklist_mobile(self):
        self.a4_worklist(390, False)

    def test_a4_full_worklist_desktop(self):
        self.a4_worklist(1366, True)

    def test_a4_full_worklist_mobile(self):
        self.a4_worklist(390, True)

    def test_a4_no_pdf_entry_has_visible_official_number(self):
        from inventory.models import Movement

        page = self.page("entry", 390)
        self.add_reference(page, self.first.name)
        page.get_by_label("Pessoa que entregou", exact=True).fill("Cliente A4 browser fictício")
        page.get_by_label("Garrafas soltas", exact=True).fill("2")
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        self.assertIn("Em preparação", page.locator("h1").inner_text())
        page.locator('select[name="report_requested"]').select_option("no")
        page.locator('input[name="confirm"]').check()
        page.get_by_role("button", name="Confirmar entrada", exact=True).click()
        movement = self.database_call(lambda: Movement.objects.latest("pk"))
        official = self.database_call(lambda: str(movement.report_number))
        self.assertIn(official, page.locator("h1").inner_text())
        self.assertIn("Operação confirmada · " + official, page.locator("body").inner_text())
        self.assertEqual(self.database_call(lambda: movement.documents.count()), 0)
        self.assertEqual(self.stocks(), [17, 5])
        self.assert_responsive(page, 390)


class Stage4BrowserTests(StaticLiveServerTestCase):
    database_call = staticmethod(BrowserSmokeTests.database_call)
    setUp = OperationalBrowserTests.setUp
    page = OperationalBrowserTests.page
    stocks = OperationalBrowserTests.stocks
    add_reference = OperationalBrowserTests.add_reference
    assert_responsive = OperationalBrowserTests.assert_responsive

    def signed_exit(self, width, *, touch=False):
        import io

        from pypdf import PdfReader

        from inventory.models import Movement

        page = self.page("exit", width)
        self.add_reference(page, "Azul sintética operacional")
        page.get_by_label("Pessoa que levantou", exact=True).fill("Pessoa assinatura sintética")
        page.get_by_label("Garrafas soltas", exact=True).fill("2")
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        page.get_by_label("Validação da saída", exact=True).select_option("signed")
        canvas = page.locator("#signature-canvas")

        def draw():
            canvas.evaluate("element => element.scrollIntoView({block: 'center'})")
            page.evaluate(
                "() => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))"
            )
            bounds = canvas.bounding_box()
            self.assertIsNotNone(bounds)
            points = [
                (bounds["x"] + bounds["width"] * x, bounds["y"] + bounds["height"] * y)
                for x, y in [(0.1, 0.7), (0.25, 0.2), (0.4, 0.8), (0.55, 0.3), (0.8, 0.65)]
            ]
            self.assertTrue(
                canvas.evaluate(
                    "(element, point) => document.elementFromPoint(point[0], point[1]) === element",
                    points[0],
                )
            )
            if touch:
                session = page.context.new_cdp_session(page)
                session.send("Emulation.setTouchEmulationEnabled", {"enabled": True})
                session.send(
                    "Input.dispatchTouchEvent",
                    {"type": "touchStart", "touchPoints": [{"x": points[0][0], "y": points[0][1]}]},
                )
                for x, y in points[1:]:
                    session.send(
                        "Input.dispatchTouchEvent",
                        {"type": "touchMove", "touchPoints": [{"x": x, "y": y}]},
                    )
                session.send("Input.dispatchTouchEvent", {"type": "touchEnd", "touchPoints": []})
                session.detach()
            else:
                page.mouse.move(*points[0])
                page.mouse.down()
                for x, y in points[1:]:
                    page.mouse.move(x, y, steps=5)
                page.mouse.up()

        draw()
        self.assertTrue(
            page.locator('input[name="signature_data"]')
            .input_value()
            .startswith("data:image/png;base64,")
        )
        page.get_by_role("button", name="Limpar assinatura", exact=True).click()
        self.assertEqual(page.locator('input[name="signature_data"]').input_value(), "")
        draw()
        page.locator('input[name="confirm"]').check()
        self.assert_responsive(page, width)
        page.get_by_role("button", name="Confirmar saída", exact=True).click()
        self.assertIn("CONFIRMADO", page.locator("body").inner_text())
        self.assertEqual(self.stocks(), [13, 5])
        link = page.get_by_role("link", name="Abrir / imprimir PDF", exact=True)
        response = page.request.get(self.live_server_url + link.get_attribute("href"))
        self.assertEqual(response.status, 200)
        self.assertIn("no-store", response.headers["cache-control"])
        pdf = PdfReader(io.BytesIO(response.body()))
        self.assertTrue(any(list(p.images) for p in pdf.pages))
        self.assertIn("Pessoa assinatura sintética", "".join(p.extract_text() for p in pdf.pages))
        self.assertEqual(
            self.database_call(lambda: Movement.objects.get(kind="exit").validation), "signed"
        )
        self.assert_responsive(page, width)

    def test_signed_exit_desktop_mouse(self):
        self.signed_exit(1366)

    def test_signed_exit_mobile_touch(self):
        self.signed_exit(390, touch=True)

    def test_pwa_manifest_installability_and_private_offline_policy(self):
        import tempfile

        directory = tempfile.TemporaryDirectory(prefix="gg-pwa-profile-")
        self.addCleanup(directory.cleanup)
        context = self.playwright.chromium.launch_persistent_context(
            directory.name,
            executable_path=os.getenv("CHROMIUM_PATH") or shutil.which("chromium"),
            headless=True,
            args=["--no-sandbox"],
            viewport={"width": 390, "height": 844},
        )
        self.addCleanup(context.close)
        page = context.new_page()
        page.set_default_timeout(10000)
        page.goto(self.live_server_url + reverse("web:entry"))
        page.get_by_label("Utilizador", exact=True).fill(self.admin.username)
        page.get_by_label("Palavra-passe", exact=True).fill(self.password)
        page.get_by_role("button", name="Entrar", exact=True).click()
        page.evaluate("async () => { await navigator.serviceWorker.ready; }")
        page.reload()
        page.wait_for_function("navigator.serviceWorker.controller !== null")
        session = page.context.new_cdp_session(page)
        manifest = session.send("Page.getAppManifest")
        self.assertEqual(manifest["errors"], [])
        self.assertEqual(session.send("Page.getInstallabilityErrors")["installabilityErrors"], [])
        session.detach()
        self.add_reference(page, "Azul sintética operacional")
        page.get_by_label("Pessoa que entregou", exact=True).fill("Pessoa offline sintética")
        page.get_by_label("Garrafas soltas", exact=True).fill("3")
        page.get_by_role("button", name="Guardar e conferir operação", exact=True).click()
        review_url = page.url
        page.context.set_offline(True)
        page.wait_for_function(
            "!navigator.onLine && document.querySelector('button[type=submit]').disabled"
        )
        result = page.evaluate(
            "async () => { try { await fetch(location.href, {method:'POST',body:new FormData(document.querySelector('form.card'))}); return 'unexpected'; } catch (_) { return 'blocked'; } }"
        )
        self.assertEqual(result, "blocked")
        page.goto(self.live_server_url + reverse("web:stock"))
        self.assertIn("Sem ligação", page.locator("h1").inner_text())
        self.assertNotIn(self.first.name, page.locator("body").inner_text())
        self.assertNotIn(self.admin.username, page.locator("body").inner_text())
        cached = page.evaluate(
            "async () => { const urls=[]; for (const name of await caches.keys()) for (const request of await (await caches.open(name)).keys()) urls.push(new URL(request.url).pathname); return urls; }"
        )
        self.assertTrue(cached)
        self.assertTrue(all(path == "/offline/" or path.startswith("/static/") for path in cached))
        self.assertEqual(self.stocks(), [15, 5])
        page.context.set_offline(False)
        page.goto(review_url)
        self.assertIn("CONFERÊNCIA", page.locator("body").inner_text())
        self.assertEqual(self.stocks(), [15, 5])

    def test_mobile_user_creation_and_excel_download(self):
        from openpyxl import load_workbook

        page = self.page("entry", 390)
        page.goto(self.live_server_url + reverse("web:user-create"))
        page.locator("#id_username").fill("browser-created-synthetic")
        page.locator("#id_role").select_option("viewer")
        password = secrets.token_urlsafe(32)
        page.locator("#id_password1").fill(password)
        page.locator("#id_password2").fill(password)
        page.locator("#id_current_password").fill(self.password)
        page.get_by_role("button", name="Guardar", exact=True).click()
        self.assertIn("browser-created-synthetic", page.locator("body").inner_text())
        self.assert_responsive(page, 390)
        page.goto(self.live_server_url + reverse("web:export"))
        page.get_by_label("Dados a exportar").select_option("references")
        with page.expect_download() as download:
            page.get_by_role("button", name="Descarregar .xlsx", exact=True).click()
        import io
        from pathlib import Path

        self.assertTrue(download.value.suggested_filename.endswith(".xlsx"))
        workbook = load_workbook(io.BytesIO(Path(download.value.path()).read_bytes()))
        self.assertEqual(workbook["Dados"].max_row, 3)
        self.assert_responsive(page, 390)


@override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"])
class PublicBrandBrowserTests(SimpleTestCase):
    """Real Chromium/public HTTP only. No accounts, stock or database setup."""

    def setUp(self):
        server = LiveServerThread("localhost", StaticFilesHandler)
        server.daemon = True
        server.start()
        self.addCleanup(server.terminate)
        self.assertTrue(server.is_ready.wait(10), "Public test server did not start")
        if server.error:
            raise server.error
        self.url = f"http://localhost:{server.port}"
        playwright = sync_playwright().start()
        self.addCleanup(playwright.stop)
        self.browser = playwright.chromium.launch(
            executable_path=os.getenv("CHROMIUM_PATH") or shutil.which("chromium"),
            headless=True,
            args=["--no-sandbox"],
        )
        self.addCleanup(self.browser.close)

    def check_login(self, width):
        import hashlib
        from pathlib import Path

        page = self.browser.new_page(viewport={"width": width, "height": 900})
        self.addCleanup(page.close)
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        response = page.goto(self.url + reverse("login"))
        self.assertEqual(response.status, 200)
        self.assertEqual(page.title(), "Entrar · Gestão Vertente Aveludada")
        logo = page.get_by_role("img", name="Copo com Vida", exact=True)
        logo.wait_for(state="visible")
        page.wait_for_function("document.querySelector('.brand-logo').naturalWidth > 0")
        dimensions = logo.evaluate(
            "image => ({natural:[image.naturalWidth,image.naturalHeight],width:image.clientWidth,height:image.clientHeight})"
        )
        self.assertEqual(dimensions["natural"], [1273, 827])
        self.assertAlmostEqual(dimensions["width"] / dimensions["height"], 1273 / 827, delta=0.02)
        image = page.request.get(self.url + logo.get_attribute("src"))
        self.assertEqual(image.status, 200)
        self.assertEqual(
            hashlib.sha256(image.body()).hexdigest(),
            "58e89b8b1c0d930e8418bae77565367f7785e55740ac435c2d3ba88555c24948",
        )
        manifest = page.request.get(self.url + reverse("manifest")).json()
        for icon in manifest["icons"]:
            self.assertIn("branding/copo-com-vida-", icon["src"])
            self.assertEqual(page.request.get(self.url + icon["src"]).status, 200)
        self.assertLessEqual(page.evaluate("document.documentElement.scrollWidth"), width)
        self.assertEqual(errors, [])
        if directory := os.getenv("BROWSER_ARTIFACT_DIR"):
            Path(directory).mkdir(parents=True, exist_ok=True)
            page.screenshot(
                path=str(Path(directory) / f"branding-login-{width}.png"), full_page=True
            )

    def test_original_logo_login_desktop(self):
        self.check_login(1366)

    def test_original_logo_login_mobile(self):
        self.check_login(390)
